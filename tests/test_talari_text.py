"""Talari's writing side, against a fake DeepSeek — no network.

DeepSeek is replaced where Talari calls it (``talari.llm._chat_completion``), so these
tests see exactly the prompts that would be sent and decide what comes back. Documents
for the import tests are made here, with python-docx and PyMuPDF.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import json
from pathlib import Path
from typing import Any, Callable

import httpx
import pytest
from fastapi.testclient import TestClient

from phansora.products.talari import llm as llm_module
from phansora.products.talari import markup
from phansora.products.talari.comfy import ComfyUI
from phansora.products.talari.config import Settings
from phansora.products.talari.schemas import IllustrationRequest
from phansora.products.talari.server import create_app
from phansora.products.talari.sources import SCANNED_PDF
from phansora.products.talari.styles import build_prompts

Reply = str | dict | Exception | Callable[[str, str], Any]


def make_settings(tmp_path: Path) -> Settings:
    return Settings(
        comfyui_url="http://comfy.test",
        comfyui_checkpoint="sdxl_base.safetensors",
        comfyui_workflow="",
        comfyui_base_size=1024,
        comfyui_steps=20,
        comfyui_cfg=7.0,
        comfyui_sampler="euler",
        comfyui_scheduler="normal",
        comfyui_timeout_s=5.0,
        data_dir=tmp_path / "talari",
        comfyui_pipeline="checkpoint",
    )


class FakeDeepSeek:
    """Stands in for the shared helper: remembers each call, answers from ``replies``.

    A reply is the text to return, a dict (returned as JSON), an exception to raise, or a
    function of (system, user) returning any of those — for calls made concurrently.
    """

    def __init__(self) -> None:
        self.calls: list[dict[str, Any]] = []
        self.replies: list[Reply] = []

    async def __call__(self, *, cfg, system_prompt, user_text, max_output_tokens, disable_thinking=False, temperature=0):
        self.calls.append(
            {"system": system_prompt, "user": user_text, "max_tokens": max_output_tokens, "temperature": temperature, "model": cfg.model}
        )
        reply: Any = self.replies.pop(0) if self.replies else ""
        if callable(reply):
            reply = reply(system_prompt, user_text)
        if isinstance(reply, Exception):
            raise reply
        return reply if isinstance(reply, str) else json.dumps(reply)


@pytest.fixture
def deepseek(monkeypatch) -> FakeDeepSeek:
    monkeypatch.setenv("DEEPSEEK_API_KEY", "test-key")
    monkeypatch.setenv("DEEPSEEK_MODEL", "deepseek-test")
    monkeypatch.delenv("TALARI_MODEL", raising=False)
    fake = FakeDeepSeek()
    monkeypatch.setattr(llm_module, "_chat_completion", fake)
    return fake


@pytest.fixture
def make_client(tmp_path: Path, deepseek: FakeDeepSeek):
    clients: list[TestClient] = []

    def build(auth: dict | None = None) -> TestClient:
        app = create_app(make_settings(tmp_path), transport=httpx.MockTransport(lambda _r: httpx.Response(404)))
        if auth is not None:
            # What the parent app's AuthGate would have put on the request.
            @app.middleware("http")
            async def as_caller(request, call_next):
                request.scope.setdefault("state", {})["auth"] = auth
                return await call_next(request)

        client = TestClient(app)
        client.__enter__()
        clients.append(client)
        return client

    yield build
    for client in clients:
        client.__exit__(None, None, None)


@pytest.fixture
def client(make_client) -> TestClient:
    return make_client()


def words(n: int, word: str = "word") -> str:
    return " ".join([word] * n)


# --- markdown to inline HTML ------------------------------------------------------


def test_model_text_becomes_inline_html_with_everything_else_escaped() -> None:
    assert markup.inline_html("A *quiet* and **loud** day & night") == "A <em>quiet</em> and <strong>loud</strong> day &amp; night"
    assert markup.inline_html("<script>alert(1)</script>") == "&lt;script&gt;alert(1)&lt;/script&gt;"
    assert markup.inline_html('<img src=x onerror="alert(1)">hi') == '&lt;img src=x onerror="alert(1)"&gt;hi'
    assert markup.inline_html("<i>tilted</i> and <b>heavy</b>") == "<em>tilted</em> and <strong>heavy</strong>"
    # Unbalanced tags are closed or dropped, never left to swallow the page.
    assert markup.inline_html("<strong>never closed") == "<strong>never closed</strong>"
    assert markup.inline_html("stray</em> close") == "stray close"
    # Not emphasis: an asterisk between spaces, and underscores inside a word.
    assert markup.inline_html("2 * 3 * 4 and snake_case_name") == "2 * 3 * 4 and snake_case_name"
    assert markup.inline_html("[a link](http://x.test) and `code`") == "a link and code"


def test_markdown_becomes_blocks() -> None:
    blocks = markup.parse_blocks(
        "## A heading\n\nFirst line\nsame paragraph.\n\n### Smaller\n\n- one\n- two *x*\n\n1. first\n2. second\n\n"
        "> Quoted\n> words\n\n**Bold Heading**\n\n* * *\n\nLast."
    )
    assert blocks == [
        {"type": "heading", "level": 2, "html": "A heading"},
        {"type": "paragraph", "html": "First line same paragraph."},
        {"type": "heading", "level": 3, "html": "Smaller"},
        {"type": "list", "items": ["one", "two <em>x</em>"]},
        {"type": "list", "items": ["first", "second"], "ordered": True},
        {"type": "quote", "html": "Quoted words"},
        {"type": "heading", "level": 3, "html": "Bold Heading"},
        {"type": "paragraph", "html": "* * *"},
        {"type": "paragraph", "html": "Last."},
    ]


def test_html_from_the_model_is_read_as_structure() -> None:
    blocks = markup.parse_blocks("<h2>Title</h2><p>One <em>two</em></p><ul><li>a</li><li>b</li></ul><p><script>x()</script></p>")
    assert blocks == [
        {"type": "heading", "level": 2, "html": "Title"},
        {"type": "paragraph", "html": "One <em>two</em>"},
        {"type": "list", "items": ["a", "b"]},
        {"type": "paragraph", "html": "&lt;script&gt;x()&lt;/script&gt;"},
    ]


# --- import -------------------------------------------------------------------------


def upload(client: TestClient, name: str, data: bytes, user: str = "7", **form: str):
    return client.post(f"/users/{user}/sources/extract", files={"file": (name, data)}, data=form)


def test_a_text_file_is_read_into_its_chapters(client: TestClient) -> None:
    text = (
        "Chapter 1\n\n"
        + words(120, "night")
        + "\n\nPage 12\n\n"
        + words(80, "storm")
        + "\n\nCHAPTER TWO\nThe Harbour\n\n"
        + words(200, "boat")
    )
    response = upload(client, "my_novel.txt", text.encode("utf-8"))
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["title"] == "my novel" and body["format"] == "txt" and body["pages"] is None
    assert [c["title"] for c in body["chapters"]] == ["Chapter 1", "CHAPTER TWO: The Harbour"]
    first, second = body["chapters"]
    assert first["index"] == 0 and first["words"] == 200 and first["pageStart"] is None and first["pageEnd"] is None
    assert "Page 12" not in first["text"]  # a page number is apparatus
    assert first["text"] == words(120, "night") + "\n\n" + words(80, "storm")
    assert second["words"] == 200 and body["words"] == 400


def test_text_without_chapters_is_cut_into_parts_at_paragraphs(client: TestClient) -> None:
    paragraphs = [words(100, f"p{n}") for n in range(65)]
    response = upload(client, "notes.md", "\n\n".join(paragraphs).encode(), title="My Notes")
    body = response.json()
    assert body["title"] == "My Notes" and body["format"] == "md"
    assert [c["title"] for c in body["chapters"]] == ["Part 1", "Part 2"]
    assert [c["words"] for c in body["chapters"]] == [3000, 3500]  # the short last part joins the one before
    assert body["chapters"][1]["text"].split("\n\n")[0] == paragraphs[30]


def test_a_docx_is_read_by_its_headings(client: TestClient) -> None:
    import docx

    document = docx.Document()
    document.add_heading("The Beginning", level=1)
    document.add_paragraph(words(50, "alpha"))
    document.add_paragraph(words(30, "beta"))
    document.add_heading("The End", level=1)
    document.add_paragraph(words(40, "omega"))
    data = io.BytesIO()
    document.save(data)
    body = upload(client, "story.docx", data.getvalue()).json()
    assert body["format"] == "docx"
    assert [(c["title"], c["words"]) for c in body["chapters"]] == [("The Beginning", 80), ("The End", 40)]


def test_a_pdf_is_read_by_its_outline_with_pages(client: TestClient) -> None:
    import fitz

    pdf = fitz.open()
    for number in range(1, 4):
        page = pdf.new_page()
        page.insert_textbox(fitz.Rect(72, 72, 520, 770), words(60, f"page{number}"), fontsize=11)
    pdf.set_toc([[1, "First Steps", 1], [1, "Further On", 3]])
    body = upload(client, "guide.pdf", pdf.tobytes()).json()
    assert body["format"] == "pdf" and body["pages"] == 3
    assert [(c["title"], c["pageStart"], c["pageEnd"]) for c in body["chapters"]] == [
        ("First Steps", 1, 2),
        ("Further On", 3, 3),
    ]
    assert "page3" in body["chapters"][1]["text"] and "page3" not in body["chapters"][0]["text"]


def test_a_scanned_pdf_is_refused_plainly(client: TestClient) -> None:
    import fitz

    pdf = fitz.open()
    for _ in range(3):
        pdf.new_page().draw_rect(fitz.Rect(50, 50, 200, 200))
    response = upload(client, "scan.pdf", pdf.tobytes())
    assert response.status_code == 422
    assert response.json()["detail"] == SCANNED_PDF


def test_unsupported_damaged_and_empty_files(client: TestClient) -> None:
    unsupported = upload(client, "book.mobi", b"not supported")
    assert unsupported.status_code == 415 and "PDF" in unsupported.json()["detail"]
    damaged = upload(client, "broken.docx", b"this is not a zip file")
    assert damaged.status_code == 422 and "could not be read" in damaged.json()["detail"]
    assert upload(client, "empty.txt", b"").status_code == 422


# --- analysis -------------------------------------------------------------------------


def test_analysis_budgets_the_text_and_returns_one_summary_per_chapter(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [
        {
            "summary": "A book about boats.",
            "themes": ["the sea", "craft"],
            "concepts": [{"name": "Keel", "description": "The spine of a boat.", "chapterIndexes": [3, 99]}],
            "audience": "Sailors.",
            "tone": "practical",
            "chapters": [{"index": 3, "summary": "Building the hull."}, {"index": 5, "summary": "Launching."}],
        }
    ]
    long_text = words(60000, "plank")  # about 360k characters
    response = client.post(
        "/users/7/sources/analyze",
        json={"kind": "book", "title": "Boats", "chapters": [{"index": 3, "title": "Hull", "text": long_text}, {"index": 5, "title": "Launch", "text": "Into the water."}]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["summary"] == "A book about boats." and body["themes"] == ["the sea", "craft"]
    assert body["concepts"] == [{"name": "Keel", "description": "The spine of a boat.", "chapterIndexes": [3]}]
    assert body["chapters"] == [
        {"index": 3, "title": "Hull", "summary": "Building the hull."},
        {"index": 5, "title": "Launch", "summary": "Launching."},
    ]
    prompt = deepseek.calls[0]["user"]
    assert len(prompt) < 90_000 and "[…]" in prompt and "Into the water." in prompt


def test_a_long_book_is_summarized_in_batches_then_as_a_whole(client: TestClient, deepseek: FakeDeepSeek) -> None:
    import re

    def answer(_system: str, user: str) -> Any:
        if "here is a summary of each chapter" in user:
            return {"summary": "A long book.", "themes": ["length"], "concepts": [], "audience": "Patient readers.", "tone": "steady"}
        asked = [int(i) for i in re.findall(r"^=== \[(\d+)\]", user, flags=re.MULTILINE)]
        # The model skips chapter 44, every time it is asked.
        return {"chapters": [{"index": i, "summary": f"Summary {i}."} for i in asked if i != 44]}

    deepseek.replies = [answer] * 10
    chapters = [{"index": i, "title": f"C{i}", "text": f"Opening words of chapter {i}. " + words(50)} for i in range(45)]
    body = client.post("/users/7/sources/analyze", json={"title": "Long", "chapters": chapters}).json()
    assert body["summary"] == "A long book."
    assert [c["index"] for c in body["chapters"]] == list(range(45))
    assert body["chapters"][3]["summary"] == "Summary 3."
    assert body["chapters"][44]["summary"].startswith("Opening words of chapter 44.")  # its own words, as a last resort
    # Two batches, one retry for the skipped chapter, one overview from the summaries.
    assert len(deepseek.calls) == 4
    assert "[3] C3: Summary 3." in deepseek.calls[-1]["user"]



def test_parts_without_chapter_titles_are_given_real_ones(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [
        {
            "summary": "A piano method.",
            "themes": ["reading music"],
            "concepts": [],
            "audience": "Adult beginners.",
            "tone": "encouraging",
            "chapters": [
                {"index": 0, "title": "“Middle C Position and Finger Numbers”", "summary": "Sitting, fingers, Middle C."},
                {"index": 1, "title": "Part 2", "summary": "The G7 chord."},
                {"index": 2, "title": "A renamed real chapter", "summary": "Waltz time."},
            ],
        }
    ]
    chapters = [
        {"index": 0, "title": "Part 1", "text": "How to Sit at the Piano. Finger Numbers. " + words(40)},
        {"index": 1, "title": "Part 2", "text": "The G7 Chord. " + words(40)},
        {"index": 2, "title": "Waltz Time", "text": "Three beats to a measure. " + words(40)},
    ]
    body = client.post("/users/7/sources/analyze", json={"title": "Alfred", "chapters": chapters}).json()
    # A placeholder takes the model's title, quotes stripped; a title the model left as a
    # placeholder stays as it was; a chapter the source named keeps the source's name.
    assert [c["title"] for c in body["chapters"]] == ["Middle C Position and Finger Numbers", "Part 2", "Waltz Time"]
    prompt = deepseek.calls[0]["user"]
    assert "only a placeholder" in prompt and '"title"' in prompt


def test_a_source_with_its_own_chapter_titles_is_not_asked_to_rename_them(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"summary": "Boats.", "themes": [], "concepts": [], "audience": "", "tone": "", "chapters": []}]
    client.post("/users/7/sources/analyze", json={"title": "Boats", "chapters": [{"index": 0, "title": "Hull", "text": "Planks."}]})
    assert "only a placeholder" not in deepseek.calls[0]["user"]


# --- planning -------------------------------------------------------------------------

PLAN = {
    "title": "Tides",
    "subtitle": "A year at sea",
    "description": "A blurb.",
    "audience": "Readers who like the sea.",
    "tone": "warm",
    "chapters": [
        {
            "title": "Chapter 1: Setting Out",
            "summary": "Leaving port.",
            "sections": [{"title": "The quay", "topics": ["ropes", "goodbyes"]}],
            "estimatedWords": 1800,
            "sources": ["S1-0", "[S1-2]", "S9-9"],
        },
        {"title": "Landfall", "summary": "Arriving.", "sections": ["Arrival"], "estimatedWords": "not a number"},
    ],
}


def test_an_idea_becomes_a_plan(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [PLAN]
    response = client.post("/users/7/blueprints", json={"mode": "idea", "brief": {"idea": "A sailing memoir", "length": "short"}})
    assert response.status_code == 200, response.text
    (plan,) = response.json()["proposals"]
    first, second = plan["chapters"]
    assert first["title"] == "Setting Out"  # numbering is the layout's job
    assert first["sections"] == [{"title": "The quay", "topics": ["ropes", "goodbyes"]}]
    assert first["sourceRefs"] == [] and first["estimatedWords"] == 1800  # an idea has no sources to cite
    assert second["sections"] == [{"title": "Arrival", "topics": []}] and second["estimatedWords"] == 1200
    assert "A sailing memoir" in deepseek.calls[0]["user"]



def test_a_book_from_documents_is_sized_by_its_material_not_a_length(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{**PLAN, "chapters": [{"title": "One", "summary": "s"}, {"title": "Two", "summary": "s"}]}]
    request = {
        "mode": "documents",
        "brief": {"idea": "A field guide", "length": "long", "targetWords": 9000},
        "sources": [
            {"id": "d1", "title": "Interviews", "chapters": [{"index": 0, "title": "Notes", "summary": "Talks.", "words": 6000}]},
            {"id": "d2", "title": "Papers", "chapters": [{"index": 0, "title": "Study", "summary": "Data.", "words": 4000}]},
        ],
    }
    response = client.post("/users/7/blueprints", json=request)
    assert response.status_code == 200, response.text
    prompt = deepseek.calls[0]["user"]
    # The material's size, not the menu's: "long" is an idea's length and is ignored here.
    assert "about 10,000 words, so the book should come to about 9,000 words in all" in prompt
    assert "[S1-0] Notes (6,000 words): Talks." in prompt and "[S2-0] Papers" not in prompt
    assert "Plan a long book" not in prompt
    # A chapter the model gave no length takes its share of the material: 9,000 over 4 chapters.
    (plan,) = response.json()["proposals"]
    assert [c["estimatedWords"] for c in plan["chapters"]] == [2250, 2250]


def test_an_idea_keeps_the_length_chosen_for_it(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [PLAN]
    client.post("/users/7/blueprints", json={"mode": "idea", "brief": {"idea": "A memoir", "length": "long", "targetWords": 9000}})
    assert "Plan a long book" in deepseek.calls[0]["user"]


def test_a_chapter_from_sources_is_told_not_to_pad(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [WRITTEN, WRITTEN]
    client.post("/users/7/chapters/write", json={**CHAPTER, "mode": "documents", "rewrite": None})
    client.post("/users/7/chapters/write", json={**CHAPTER, "mode": "idea", "rewrite": None, "sourceExcerpts": []})
    from_sources, from_idea = (call["user"] for call in deepseek.calls)
    assert "let the material set the length" in from_sources and "about 2000 words" in from_sources
    assert "let the material set the length" not in from_idea and "about 2000 words" in from_idea

def test_several_versions_are_planned_from_distinct_directions(client: TestClient, deepseek: FakeDeepSeek) -> None:
    def answer(_system: str, user: str) -> Any:
        if "Propose them" in user:
            return {"directions": [{"title": f"Version {n}", "angle": f"Angle {n}."} for n in (1, 2, 3)]}
        mine = next(n for n in (1, 2, 3) if f"Plan this version of the book:\nVersion {n}" in user)
        return {**PLAN, "title": f"Tides {mine}"}

    deepseek.replies = [answer] * 4
    source = {
        "id": "book-abc",
        "title": "The Old Sea Book",
        "summary": "An old book.",
        "chapters": [{"index": 0, "title": "Ports", "summary": "Harbours."}, {"index": 2, "title": "Storms", "summary": "Weather."}],
    }
    response = client.post(
        "/users/7/blueprints",
        json={"mode": "rewrite", "brief": {"audience": "teenagers"}, "rewrite": {"depth": "deep", "options": ["simplify", "add-humor"]}, "sources": [source], "count": 3},
    )
    assert response.status_code == 200, response.text
    proposals = response.json()["proposals"]
    assert sorted(p["title"] for p in proposals) == ["Tides 1", "Tides 2", "Tides 3"]
    assert proposals[0]["chapters"][0]["sourceRefs"] == [
        {"sourceId": "book-abc", "chapterIndex": 0},
        {"sourceId": "book-abc", "chapterIndex": 2},
    ]
    plans = [c["user"] for c in deepseek.calls if "Propose them" not in c["user"]]
    assert len(plans) == 3
    version_1 = next(p for p in plans if "Plan this version of the book:\nVersion 1" in p)
    assert "Version 2: Angle 2." in version_1 and "Version 3: Angle 3." in version_1  # told what to differ from
    assert "Deep: a distinct new book" in version_1 and "simpler and more accessible" in version_1 and "add humor" in version_1
    assert "[S1-2] Storms: Weather." in version_1


def test_a_version_that_fails_is_left_out(client: TestClient, deepseek: FakeDeepSeek) -> None:
    def answer(_system: str, user: str) -> Any:
        if "Propose them" in user:
            return {"directions": [{"title": "A", "angle": "First."}, {"title": "B", "angle": "Second."}]}
        if "Plan this version of the book:\nB" in user:
            return RuntimeError("DeepSeek chat HTTP 500: upstream exploded")
        return PLAN

    deepseek.replies = [answer] * 10
    response = client.post("/users/7/blueprints", json={"mode": "idea", "brief": {"idea": "A memoir"}, "count": 2})
    assert response.status_code == 200 and [p["title"] for p in response.json()["proposals"]] == ["Tides"]


def test_refining_a_plan_keeps_each_chapters_sources(client: TestClient, deepseek: FakeDeepSeek) -> None:
    original = {
        "title": "Tides",
        "chapters": [
            {"title": "Ports", "sourceRefs": [{"sourceId": "a", "chapterIndex": 0}]},
            {"title": "Storms", "sourceRefs": [{"sourceId": "a", "chapterIndex": 1}]},
            {"title": "Home", "sourceRefs": [{"sourceId": "b", "chapterIndex": 4}]},
        ],
    }
    deepseek.replies = [
        {
            "reply": "I merged the first and last chapters and added one on navigation.",
            "blueprint": {
                "title": "Tides",
                "chapters": [
                    {"from": 2, "title": "Storms, Practically", "summary": "More practical."},
                    {"from": None, "title": "Navigation"},
                    {"from": [1, 3], "title": "Ports and Home"},
                ],
            },
        }
    ]
    response = client.post(
        "/users/7/blueprints/refine",
        json={"blueprint": original, "message": "Make chapter 2 more practical", "history": [{"role": "user", "content": "Hi"}, {"role": "assistant", "content": "Hello"}]},
    )
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["reply"].startswith("I merged")
    assert [(c["title"], c["sourceRefs"]) for c in body["blueprint"]["chapters"]] == [
        ("Storms, Practically", [{"sourceId": "a", "chapterIndex": 1}]),
        ("Navigation", []),
        ("Ports and Home", [{"sourceId": "a", "chapterIndex": 0}, {"sourceId": "b", "chapterIndex": 4}]),
    ]
    prompt = deepseek.calls[0]["user"]
    assert "2. Storms" in prompt and "Author: Hi" in prompt and "Make chapter 2 more practical" in prompt

    # A question changes nothing.
    deepseek.replies = [{"reply": "It has three chapters.", "blueprint": None}]
    body = client.post("/users/7/blueprints/refine", json={"blueprint": original, "message": "How many chapters?"}).json()
    assert body["reply"] == "It has three chapters." and [c["title"] for c in body["blueprint"]["chapters"]] == ["Ports", "Storms", "Home"]


# --- writing --------------------------------------------------------------------------

CHAPTER = {
    "book": {"title": "Tides", "tone": "warm", "audience": "Readers"},
    "outline": [{"title": "Setting Out"}, {"title": "Storms"}, {"title": "Landfall", "summary": "Arriving."}],
    "chapterIndex": 2,
    "chapter": {"title": "Landfall", "summary": "Arriving at last.", "sections": [{"title": "First sight", "topics": ["the coast"]}], "estimatedWords": 2000},
    "previousSummary": "The storm passed and the crew counted their losses.",
    "sourceExcerpts": [{"title": "Old log", "text": "<p>We saw land at <em>dawn</em>.</p>"}],
    "mode": "rewrite",
    "rewrite": {"depth": "light", "options": []},
    "illustrations": {"enabled": True, "perChapter": 2},
}

WRITTEN = """# Landfall

## First sight

The coast rose out of the *mist*, and **everyone** cheered.
[[ILLUSTRATION: A sailing ship approaching a misty green coast at dawn | Land at last]]

<script>alert("x")</script>

- ropes coiled
- sails furled

> Land ho!

The anchor went down.

[[ILLUSTRATION: Sailors rowing a small boat toward a beach | Going ashore]]

===SUMMARY===
The crew reached land at dawn and anchored.
"""


def test_a_chapter_is_written_as_blocks_with_pictures_and_a_recap(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [WRITTEN]
    response = client.post("/users/7/chapters/write", json=CHAPTER)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["blocks"] == [
        {"type": "heading", "level": 2, "html": "First sight"},  # the title itself is the layout's
        {"type": "paragraph", "html": "The coast rose out of the <em>mist</em>, and <strong>everyone</strong> cheered."},
        {"type": "paragraph", "html": '&lt;script&gt;alert("x")&lt;/script&gt;'},
        {"type": "list", "items": ["ropes coiled", "sails furled"]},
        {"type": "quote", "html": "Land ho!"},
        {"type": "paragraph", "html": "The anchor went down."},
    ]
    assert body["illustrations"] == [
        {"afterBlock": 1, "prompt": "A sailing ship approaching a misty green coast at dawn", "caption": "Land at last"},
        {"afterBlock": 5, "prompt": "Sailors rowing a small boat toward a beach", "caption": "Going ashore"},
    ]
    assert body["summary"] == "The crew reached land at dawn and anchored."

    (call,) = deepseek.calls
    prompt = call["user"]
    assert "The storm passed and the crew counted their losses." in prompt  # continuity
    assert "Write chapter 3, “Landfall”" in prompt and "<- this chapter" in prompt
    assert "We saw land at *dawn*." in prompt and "Keep its structure and everything it says" in prompt
    assert "about 2000 words" in prompt and "Mark 2 places for a picture" in prompt
    assert call["max_tokens"] >= 3200 and call["temperature"] < 1


def test_a_cut_off_chapter_ends_on_its_last_whole_paragraph_and_gets_a_recap(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = ["The first paragraph is whole.\n\nThe second one stops in the mid", "A recap written separately."]
    body = client.post("/users/7/chapters/write", json={**CHAPTER, "illustrations": {"enabled": False, "perChapter": 0}}).json()
    assert body["blocks"] == [{"type": "paragraph", "html": "The first paragraph is whole."}]
    assert body["summary"] == "A recap written separately." and body["illustrations"] == []
    assert "Mark " not in deepseek.calls[0]["user"]


def test_a_section_headed_summary_stays_in_the_chapter(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = ["Body.\n\n## Summary\n\nThe key points.\n\n**===SUMMARY===**\nThe recap."]
    body = client.post("/users/7/chapters/write", json={**CHAPTER, "illustrations": {"enabled": False}}).json()
    assert [b.get("html") for b in body["blocks"]] == ["Body.", "Summary", "The key points."]
    assert body["summary"] == "The recap."


def test_pictures_the_chapter_did_not_place_are_chosen_from_its_text(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [
        "One calm paragraph.\n\nThe gulls wheeled over the white cliffs.\n\n===SUMMARY===\nCalm.",
        {"prompts": [{"prompt": "Gulls over white cliffs", "caption": "Gulls", "anchor": "the gulls wheeled over the white cliffs"}]},
    ]
    body = client.post("/users/7/chapters/write", json={**CHAPTER, "illustrations": {"enabled": True, "perChapter": 1}}).json()
    assert body["illustrations"] == [{"afterBlock": 1, "prompt": "Gulls over white cliffs", "caption": "Gulls"}]


# --- revising ---------------------------------------------------------------------------


def revise(client: TestClient, **fields: Any):
    payload = {"action": "rewrite", "html": "<p>The old text.</p>", "context": {"bookTitle": "Tides", "chapterTitle": "Storms"}}
    return client.post("/users/7/text/revise", json={**payload, **fields})


def test_a_revision_comes_back_as_inline_html_with_notes(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = ["New <b>bold</b> text with <script>x()</script> & *style*.\n\n===NOTES===\n- Check the date given."]
    response = revise(client, instruction="Keep it short", context={"before": "<p>Earlier.</p>", "after": "Later."})
    assert response.status_code == 200, response.text
    assert response.json() == {
        "html": "<p>New <strong>bold</strong> text with &lt;script&gt;x()&lt;/script&gt; &amp; <em>style</em>.</p>",
        "notes": ["Check the date given."],
    }
    prompt = deepseek.calls[0]["user"]
    assert "The old text." in prompt and "Earlier." in prompt and "Later." in prompt and "Keep it short" in prompt


def test_one_plain_paragraph_comes_back_unwrapped_and_several_come_back_as_paragraphs(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = ["Simpler.", "One.\n\nTwo."]
    assert revise(client, action="simplify", html="Hard words.").json()["html"] == "Simpler."
    assert revise(client, action="expand", html="Short.").json()["html"] == "<p>One.</p><p>Two.</p>"


def test_continue_keeps_the_passage_and_adds_to_it(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = ["And then the wind rose."]
    body = revise(client, action="continue", html="<p>It was <em>calm</em>.</p>").json()
    assert body["html"] == "<p>It was <em>calm</em>.</p><p>And then the wind rose.</p>"


def test_factcheck_leaves_the_text_alone_and_never_claims_a_check(client: TestClient, deepseek: FakeDeepSeek) -> None:
    html = "<p>The library opened in <em>1602</em>.</p><p>It holds 2 million books &amp; maps.</p>"
    deepseek.replies = [{"claims": ["Check the opening year (1602).", "The book count is verified as accurate."]}]
    body = revise(client, action="factcheck", html=html).json()
    assert body == {"html": html, "notes": ["Check the opening year (1602)."]}
    assert "must not judge whether anything" in deepseek.calls[0]["user"]

    deepseek.replies = [{"claims": []}]
    notes = revise(client, action="factcheck", html=html).json()["notes"]
    assert len(notes) == 1 and "Nothing here has been verified" in notes[0]


def test_revisions_that_need_more_say_so(client: TestClient, deepseek: FakeDeepSeek) -> None:
    assert revise(client, action="tone").status_code == 422
    assert revise(client, action="custom").status_code == 422
    assert revise(client, html="").status_code == 422
    assert revise(client, action="summarize").status_code == 422
    assert deepseek.calls == []


# --- illustration prompts ---------------------------------------------------------------

CHAPTER_TEXT = "<p>The keeper climbed the stairs.</p><p>At the top, the great lamp turned slowly over the black water.</p>"


def test_interior_ideas_quote_the_chapter_word_for_word(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [
        {
            "prompts": [
                {"prompt": "A lighthouse keeper beside a huge turning lamp at night", "caption": "The lamp", "anchor": "the Great Lamp turned slowly over the dark water"},
                {"prompt": "Spiral stairs inside a lighthouse", "caption": "Up", "anchor": "The keeper climbed the stairs."},
            ]
        }
    ]
    response = client.post(
        "/users/7/illustration-prompts",
        json={"book": {"title": "The Light"}, "chapter": {"title": "Night", "text": CHAPTER_TEXT}, "styleId": "custom", "customStyle": "linocut, two colours", "count": 2},
    )
    assert response.status_code == 200, response.text
    prompts = response.json()["prompts"]
    # Misquoted by the model; given back as the chapter's own words.
    assert prompts[0]["anchor"] == "the great lamp turned slowly over the"
    assert prompts[1]["anchor"] == "The keeper climbed the stairs"
    assert "linocut, two colours" in deepseek.calls[0]["user"]


def test_cover_ideas_have_no_anchor_and_no_lettering(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"prompts": [{"prompt": 'A lighthouse at dusk with a banner reading "The Light"', "caption": "Beacon", "anchor": "x"}]}]
    response = client.post("/users/7/illustration-prompts", json={"book": {"title": "The Light"}, "purpose": "cover", "styleId": "ink"})
    assert response.json() == {"prompts": [{"prompt": "A lighthouse at dusk with a banner reading", "caption": "Beacon", "anchor": ""}]}
    assert "never mention the title" in deepseek.calls[0]["user"]


# --- errors and scope ---------------------------------------------------------------------


def test_deepseek_not_configured_is_a_503_naming_the_variable(client: TestClient, monkeypatch) -> None:
    monkeypatch.delenv("DEEPSEEK_API_KEY")
    monkeypatch.delenv("DEEPSEEK_CHAT_API_KEY", raising=False)
    response = revise(client)
    assert response.status_code == 503 and "DEEPSEEK_API_KEY" in response.json()["detail"]

    monkeypatch.setenv("DEEPSEEK_API_KEY", "test-key")
    monkeypatch.delenv("DEEPSEEK_MODEL")
    monkeypatch.delenv("DEEPSEEK_CHAT_MODEL", raising=False)
    response = revise(client)
    assert response.status_code == 503 and "TALARI_MODEL" in response.json()["detail"]


def test_upstream_failures_are_502s_without_the_providers_body(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [RuntimeError('DeepSeek clean failed after retries: DeepSeek chat HTTP 402: {"error": "secret internals"}')]
    response = revise(client)
    assert response.status_code == 502
    assert "out of credit" in response.json()["detail"] and "secret" not in response.json()["detail"]

    deepseek.replies = ["not json at all", "still not json"]
    response = client.post("/users/7/blueprints", json={"mode": "idea", "brief": {"idea": "x"}})
    assert response.status_code == 502 and "could not be read" in response.json()["detail"]
    assert len(deepseek.calls) == 3  # the unreadable reply was asked for once more


def test_a_truncated_json_reply_keeps_its_complete_part(client: TestClient, deepseek: FakeDeepSeek) -> None:
    whole = json.dumps({"title": "Tides", "chapters": [{"title": "One"}, {"title": "Two"}]})
    deepseek.replies = [whole[: whole.index('{"title": "Two"') + 12]]
    body = client.post("/users/7/blueprints", json={"mode": "idea", "brief": {"idea": "x"}}).json()
    assert [c["title"] for c in body["proposals"][0]["chapters"]] == ["One"]


def test_requests_are_validated(client: TestClient, deepseek: FakeDeepSeek) -> None:
    assert client.post("/users/7/blueprints", json={"mode": "idea", "brief": {"idea": "x"}, "count": 4}).status_code == 422
    assert client.post("/users/7/blueprints", json={"mode": "rewrite", "brief": {"idea": "x"}}).status_code == 422  # no sources
    assert client.post("/users/7/sources/analyze", json={"chapters": []}).status_code == 422
    assert client.post("/users/7/chapters/write", json={"book": {}}).status_code == 422
    assert client.post("/users/7/chapters/write", json={"chapter": None}).status_code == 422
    assert deepseek.calls == []

    # A null is a field left out, not an error.
    deepseek.replies = ["Written.\n\n===SUMMARY===\nDone."]
    nulls = {"book": {"title": "T", "subtitle": None}, "chapter": {"title": "One", "estimatedWords": None, "sections": None}, "previousSummary": None, "rewrite": None, "illustrations": None}
    response = client.post("/users/7/chapters/write", json=nulls)
    assert response.status_code == 200, response.text
    assert "about 1500 words" in deepseek.calls[0]["user"]


def test_a_browser_token_only_reaches_its_own_user(make_client, deepseek: FakeDeepSeek) -> None:
    client = make_client(auth={"internal": False, "user_id": "7"})
    deepseek.replies = ["Fine."]
    assert revise(client).status_code == 200
    assert client.post("/users/8/text/revise", json={"action": "rewrite", "html": "x"}).status_code == 403
    assert client.post("/users/8/blueprints", json={"mode": "idea", "brief": {"idea": "x"}}).status_code == 403
    assert upload(client, "a.txt", b"Some text.", user="8").status_code == 403
    assert upload(client, "a.txt", b"Some text.", user="7").status_code == 200
    assert len(deepseek.calls) == 1


# --- styles -------------------------------------------------------------------------------


def test_custom_styles_palette_and_detail_reach_the_picture(tmp_path: Path) -> None:
    unchanged, _ = build_prompts("A moonlit tower.", "ink", "interior")
    assert unchanged == "A moonlit tower, pen and ink illustration, confident black line work, cross-hatching, high contrast, white paper"

    positive, negative = build_prompts("A fox", "custom", "interior", custom_style='  "woodcut", bold  ', palette="grayscale", detail="rich")
    assert positive == (
        "A fox, woodcut, bold, black and white, grayscale, monochrome, "
        "richly detailed, intricate textures, layered background detail"
    )
    assert "colour" in negative and "lettering" in negative
    assert build_prompts("A fox", "custom", "interior")[0] == build_prompts("A fox", "watercolor", "interior")[0]
    # The look as the app now picks it: an art style, then its medium; "auto" adds no medium.
    # A known art style wins over the older style id, which an unknown one falls back to.
    positive, _ = build_prompts("A fox", "watercolor", "interior", art_style="realism", medium="oil-painting")
    assert positive.startswith("A fox, photorealistic illustration") and "oil painting, rich impasto" in positive
    assert "watercolor illustration" not in positive
    auto, _ = build_prompts("A fox", "", "interior", art_style="anime", medium="auto")
    assert auto == "A fox, anime and manga style, expressive large eyes, clean cel-shaded line art, vibrant colors"
    assert build_prompts("A fox", "ink", "interior", art_style="cubism")[0] == build_prompts("A fox", "ink", "interior")[0]
    # "none": the prompt as written, nothing added — what the styleguide's ComfyUI page sends.
    assert build_prompts("A fox in the snow", "none", "interior")[0] == "A fox in the snow"
    for style in ("childrens", "lineart", "comic", "vintage", "minimal", "digital"):
        assert build_prompts("A fox", style, "interior")[0] != build_prompts("A fox", "watercolor", "interior")[0]

    request = IllustrationRequest.model_validate(
        {"prompt": "A fox", "styleId": "custom", "aspect": 1, "customStyle": "paper cut-out collage", "palette": "grayscale", "detail": "simple"}
    )

    async def workflow() -> dict:
        async with httpx.AsyncClient(transport=httpx.MockTransport(lambda _r: httpx.Response(404))) as http:
            return await ComfyUI(make_settings(tmp_path), http)._workflow(request, seed=1)

    text = asyncio.run(workflow())["2"]["inputs"]["text"]
    assert text.startswith("A fox, paper cut-out collage, black and white") and "simple composition" in text
    with pytest.raises(ValueError):
        IllustrationRequest.model_validate({"prompt": "A fox", "styleId": "custom", "aspect": 1, "customStyle": "x" * 401})


# --- what an import holds besides its text (material.py, library.py) ---------------------


def png(width: int = 320, height: int = 220, color: tuple[int, int, int] = (200, 40, 40)) -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (width, height), color).save(out, "PNG")
    return out.getvalue()


def test_a_pdf_without_an_outline_is_read_by_the_size_of_its_type(client: TestClient) -> None:
    import fitz

    pdf = fitz.open()
    for number, title in enumerate(["Getting Started", "Going Further"], start=1):
        page = pdf.new_page()
        page.insert_textbox(fitz.Rect(72, 60, 520, 100), title, fontsize=22)
        page.insert_textbox(fitz.Rect(72, 110, 520, 330), words(120, f"intro{number}"), fontsize=11)
        page.insert_textbox(fitz.Rect(72, 350, 520, 380), f"A Closer Look {number}", fontsize=15)
        page.insert_textbox(fitz.Rect(72, 390, 520, 700), words(140, f"detail{number}"), fontsize=11)
    body = upload(client, "handbook.pdf", pdf.tobytes()).json()
    assert [c["title"] for c in body["chapters"]] == ["Getting Started", "Going Further"]
    assert "Part 1" not in [c["title"] for c in body["chapters"]]
    sections = [(s["chapterIndex"], s["title"], s["pageStart"]) for s in body["sections"]]
    assert sections == [(0, "A Closer Look 1", 1), (1, "A Closer Look 2", 2)]
    assert body["sections"][0]["text"].startswith("detail1") and "intro1" not in body["sections"][0]["text"]


def atlas_pdf() -> bytes:
    """A page with prose, a picture and its caption, and a ruled table; a page with a diagram."""
    import fitz

    pdf = fitz.open()
    page = pdf.new_page()
    page.insert_textbox(fitz.Rect(72, 40, 520, 200), words(150, "prose"), fontsize=11)
    page.insert_image(fitz.Rect(72, 210, 312, 375), stream=png())
    page.insert_text((72, 392), "Figure 1: A red field", fontsize=9)
    # A table: a ruled grid with a word in each cell.
    for row in range(4):
        page.draw_line((72, 420 + row * 20), (372, 420 + row * 20))
    for col in range(4):
        page.draw_line((72 + col * 100, 420), (72 + col * 100, 480))
    for row, cells in enumerate([["Name", "Age", "Town"], ["Ada", "36", "Bath"], ["Ben", "41", "York"]]):
        for col, cell in enumerate(cells):
            page.insert_text((78 + col * 100, 434 + row * 20), cell, fontsize=9)
    second = pdf.new_page()
    second.insert_textbox(fitz.Rect(72, 40, 520, 200), words(150, "more"), fontsize=11)
    # A diagram: three rows of touching circles, joined top to bottom.
    for n in range(15):
        second.draw_circle((150 + 56 * (n % 5), 300 + 56 * (n // 5)), 26)
        second.draw_line((150 + 56 * (n % 5), 300), (150 + 56 * (n % 5), 412))
    return pdf.tobytes()


def test_a_pdf_gives_up_its_pictures_tables_and_drawings(client: TestClient, tmp_path: Path) -> None:
    body = upload(client, "atlas.pdf", atlas_pdf()).json()
    kinds = sorted(a["kind"] for a in body["assets"])
    assert kinds == ["drawing", "picture", "table"]
    picture = next(a for a in body["assets"] if a["kind"] == "picture")
    assert picture["page"] == 1 and picture["caption"] == "Figure 1: A red field"
    assert picture["src"].startswith("/talari/users/7/images/") and (picture["width"], picture["height"]) == (320, 220)
    assert (tmp_path / "talari" / "images" / "7" / picture["src"].rsplit("/", 1)[1]).is_file()
    table = next(a for a in body["assets"] if a["kind"] == "table")
    assert table["rows"][0] == ["Name", "Age", "Town"] and table["rows"][2] == ["Ben", "41", "York"] and table["src"]
    assert next(a for a in body["assets"] if a["kind"] == "drawing")["page"] == 2


def test_an_upload_can_leave_a_pdfs_pictures_for_later(client: TestClient) -> None:
    quick = upload(client, "atlas.pdf", atlas_pdf(), assets="false").json()
    assert quick["assets"] == [] and quick["assetsPending"] is True and quick["chapters"]
    later = client.post(f"/users/7/sources/{quick['original']}/assets")
    assert later.status_code == 200
    body = later.json()
    assert sorted(a["kind"] for a in body["assets"]) == ["drawing", "picture", "table"] and body["assetsPending"] is False
    # Read the same both times, so the numbers of one reading hold for the other.
    assert body["chapters"] == quick["chapters"] and body["sections"] == quick["sections"]
    assert client.post(f"/users/7/sources/{'0' * 64}.pdf/assets").status_code == 404
    assert client.post("/users/7/sources/notes.txt/assets").status_code == 404


def test_an_upload_is_kept_at_once_and_read_later(client: TestClient, tmp_path: Path) -> None:
    data = atlas_pdf()
    kept = client.post("/users/7/sources/keep", files={"file": ("Field Atlas.pdf", data)})
    assert kept.status_code == 200
    body = kept.json()
    assert body == {"original": f"{hashlib.sha256(data).hexdigest()}.pdf", "format": "pdf", "bytes": len(data)}
    assert (tmp_path / "talari" / "sources" / "7" / body["original"]).read_bytes() == data
    quick = client.post(f"/users/7/sources/{body['original']}/read", json={"assets": False, "title": "Field Atlas"}).json()
    assert quick["title"] == "Field Atlas" and quick["chapters"] and quick["assets"] == [] and quick["assetsPending"] is True
    full = client.post(f"/users/7/sources/{body['original']}/read", json={"assets": True}).json()
    assert sorted(a["kind"] for a in full["assets"]) == ["drawing", "picture", "table"] and full["chapters"] == quick["chapters"]
    assert client.post("/users/7/sources/keep", files={"file": ("run.exe", b"MZ")}).status_code == 415
    assert client.post("/users/7/sources/keep", files={"file": ("empty.pdf", b"")}).status_code == 422
    assert client.post(f"/users/7/sources/{'0' * 64}.pdf/read", json={}).status_code == 404


def test_a_kept_document_of_any_kind_can_be_read(client: TestClient) -> None:
    from docx import Document

    doc = Document()
    doc.add_heading("Soil", level=1)
    doc.add_paragraph(words(80, "soil"))
    doc.add_heading("Water", level=1)
    doc.add_paragraph(words(80, "water"))
    out = io.BytesIO()
    doc.save(out)
    kept = client.post("/users/7/sources/keep", files={"file": ("garden.docx", out.getvalue())}).json()
    body = client.post(f"/users/7/sources/{kept['original']}/read", json={"assets": True}).json()
    assert [c["title"] for c in body["chapters"]] == ["Soil", "Water"] and body["format"] == "docx"


def test_a_kept_scan_without_text_says_so_when_it_is_read(client: TestClient) -> None:
    import fitz

    pdf = fitz.open()
    for _ in range(3):
        pdf.new_page().draw_rect(fitz.Rect(50, 50, 200, 200))
    kept = client.post("/users/7/sources/keep", files={"file": ("scan.pdf", pdf.tobytes())}).json()
    response = client.post(f"/users/7/sources/{kept['original']}/read", json={"assets": False})
    assert response.status_code == 422 and response.json()["detail"] == SCANNED_PDF


def test_files_no_longer_used_are_forgotten_and_nothing_else(client: TestClient, tmp_path: Path) -> None:
    body = upload(client, "atlas.pdf", atlas_pdf()).json()
    images = tmp_path / "talari" / "images" / "7"
    kept = tmp_path / "talari" / "sources" / "7"
    names = [a["src"].rsplit("/", 1)[1] for a in body["assets"] if a["src"]]
    assert names and all((images / n).is_file() for n in names)
    assert client.get(f"/users/7/sources/{body['original']}/pages/1").status_code == 200
    # Another user's copy of the same picture is theirs.
    (tmp_path / "talari" / "images" / "8").mkdir(parents=True)
    (tmp_path / "talari" / "images" / "8" / names[0]).write_bytes(b"theirs")
    other = images / ("f" * 32 + ".png")
    other.write_bytes(b"still used")
    response = client.post("/users/7/forget", json={
        "images": [*names, "../8/" + names[0], "not-ours.png"],
        "sources": [body["original"], "../../secret.pdf"],
    })
    assert response.status_code == 200 and response.json() == {"images": len(names), "sources": 1}
    assert not any((images / n).exists() for n in names) and other.is_file()
    assert not (kept / body["original"]).exists() and not list((kept / "pages").glob("*.png"))
    assert (tmp_path / "talari" / "images" / "8" / names[0]).read_bytes() == b"theirs"


def test_noise_from_a_text_layer_is_not_a_heading() -> None:
    from phansora.products.talari.material import _reads_as_heading

    assert _reads_as_heading("Preliminary Exercises") and _reads_as_heading("ODE TO JOY")
    for noise in ["No. 1", "|e eo ds", "8", "(al", "a", "— ——", "<d ~d -@ ~fj"]:
        assert not _reads_as_heading(noise), noise


def test_a_scanned_pdf_with_a_text_layer_is_read_in_parts_without_its_page_pictures(client: TestClient, monkeypatch) -> None:
    import fitz

    # Naming a picture's xref hashes it: minutes for a 164-page scan, so a scan never does.
    def hashes(*_args, **_kwargs):
        raise AssertionError("get_image_rects on a scan")

    monkeypatch.setattr(fitz.Page, "get_image_rects", hashes)
    pdf = fitz.open()
    for number in range(6):
        page = pdf.new_page()
        page.insert_image(page.rect, stream=png(612, 792, (236, 232, 224)))
        page.insert_textbox(fitz.Rect(72, 60, 520, 100), f"Lesson Heading {number}", fontsize=22)
        page.insert_textbox(fitz.Rect(72, 110, 520, 700), words(220, f"scan{number}"), fontsize=11)
    body = upload(client, "scan.pdf", pdf.tobytes()).json()
    assert all(c["title"].startswith("Part ") for c in body["chapters"])
    assert body["sections"] == [] and body["assets"] == []
    later = client.post(f"/users/7/sources/{body['original']}/assets").json()
    assert later["assets"] == []


def test_type_that_cuts_a_pdf_into_scraps_is_not_its_structure(client: TestClient) -> None:
    import fitz

    pdf = fitz.open()
    for number in range(4):
        page = pdf.new_page()
        for k in range(8):
            y = 50 + k * 90
            page.insert_textbox(fitz.Rect(72, y, 520, y + 32), f"Big Line {number} {k}", fontsize=20)
            page.insert_textbox(fitz.Rect(72, y + 34, 520, y + 84), words(12, f"bit{number}{k}"), fontsize=11)
    body = upload(client, "method.pdf", pdf.tobytes()).json()
    assert all(c["title"].startswith("Part ") for c in body["chapters"])
    assert body["sections"] == []


def test_chapters_sharing_a_page_do_not_share_its_sections(client: TestClient) -> None:
    import fitz

    pdf = fitz.open()
    first = pdf.new_page()
    first.insert_textbox(fitz.Rect(72, 40, 520, 85), "Getting Started", fontsize=22)
    first.insert_textbox(fitz.Rect(72, 90, 520, 330), words(160, "intro"), fontsize=11)
    first.insert_textbox(fitz.Rect(72, 335, 520, 370), "A Closer Look", fontsize=15)
    first.insert_textbox(fitz.Rect(72, 375, 520, 470), words(60, "closer"), fontsize=11)
    first.insert_textbox(fitz.Rect(72, 480, 520, 525), "Going Further", fontsize=22)
    first.insert_textbox(fitz.Rect(72, 530, 520, 760), words(140, "further"), fontsize=11)
    second = pdf.new_page()
    second.insert_textbox(fitz.Rect(72, 50, 520, 300), words(160, "more"), fontsize=11)
    second.insert_textbox(fitz.Rect(72, 315, 520, 350), "The Last Step", fontsize=15)
    second.insert_textbox(fitz.Rect(72, 355, 520, 500), words(60, "last"), fontsize=11)
    body = upload(client, "guide.pdf", pdf.tobytes()).json()
    assert [c["title"] for c in body["chapters"]] == ["Getting Started", "Going Further"]
    assert [(s["chapterIndex"], s["title"]) for s in body["sections"]] == [(0, "A Closer Look"), (1, "The Last Step")]
    closer = body["sections"][0]["text"]
    assert closer.startswith("closer") and "further" not in closer


def test_a_docx_keeps_its_heading_levels_its_tables_and_its_pictures(client: TestClient) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Soil", level=1)
    document.add_paragraph(words(40, "earth"))
    document.add_heading("Testing the Soil", level=2)
    document.add_paragraph(words(30, "acid"))
    table = document.add_table(rows=2, cols=2)
    for r, cells in enumerate([["pH", "Crop"], ["6.5", "Beans"]]):
        for c, text in enumerate(cells):
            table.cell(r, c).text = text
    document.add_picture(io.BytesIO(png()))
    document.add_heading("Water", level=1)
    document.add_paragraph(words(50, "rain"))
    data = io.BytesIO()
    document.save(data)
    body = upload(client, "garden.docx", data.getvalue()).json()
    assert [(c["title"], c["words"]) for c in body["chapters"]] == [("Soil", 70), ("Water", 50)]
    assert [(s["chapterIndex"], s["title"], s["words"]) for s in body["sections"]] == [(0, "Testing the Soil", 30)]
    assert sorted((a["kind"], a["chapterIndex"]) for a in body["assets"]) == [("picture", 0), ("table", 0)]
    assert next(a for a in body["assets"] if a["kind"] == "table")["rows"] == [["pH", "Crop"], ["6.5", "Beans"]]


def test_an_epub_gives_up_its_sections_pictures_and_tables(client: TestClient, tmp_path: Path) -> None:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_title("Field Notes")
    book.add_item(epub.EpubItem(uid="fig", file_name="images/fig.png", media_type="image/png", content=png()))
    chapters = []
    for n, name in enumerate(["Birds", "Bees"], start=1):
        chapter = epub.EpubHtml(title=name, file_name=f"ch{n}.xhtml")
        chapter.content = (
            f"<html><body><h1>{name}</h1><p>{words(60, name.lower())}</p><h2>Watching {name}</h2><p>{words(30, 'watch')}</p>"
            + ('<figure><img src="images/fig.png" alt="red"/><figcaption>Plate 1</figcaption></figure>'
               "<table><tr><th>Kind</th><th>Seen</th></tr><tr><td>Wren</td><td>3</td></tr></table>" if n == 1 else "")
            + "</body></html>"
        )
        book.add_item(chapter)
        chapters.append(chapter)
    book.spine = chapters
    book.add_item(epub.EpubNcx())
    path = tmp_path / "notes.epub"
    epub.write_epub(str(path), book)
    body = upload(client, "notes.epub", path.read_bytes()).json()
    assert [c["title"] for c in body["chapters"]] == ["Birds", "Bees"]
    assert [(s["chapterIndex"], s["title"]) for s in body["sections"]] == [(0, "Watching Birds"), (1, "Watching Bees")]
    picture = next(a for a in body["assets"] if a["kind"] == "picture")
    assert picture["caption"] == "Plate 1" and picture["chapterIndex"] == 0
    assert next(a for a in body["assets"] if a["kind"] == "table")["rows"] == [["Kind", "Seen"], ["Wren", "3"]]


def test_an_upload_is_kept_and_its_pages_can_be_looked_at_again(client: TestClient, tmp_path: Path) -> None:
    import fitz

    pdf = fitz.open()
    for number in range(2):
        pdf.new_page().insert_textbox(fitz.Rect(72, 72, 520, 770), words(80, f"page{number}"), fontsize=11)
    data = pdf.tobytes()
    body = upload(client, "pages.pdf", data).json()
    assert body["original"] == f"{hashlib.sha256(data).hexdigest()}.pdf"
    assert (tmp_path / "talari" / "sources" / "7" / body["original"]).read_bytes() == data
    page = client.get(f"/users/7/sources/{body['original']}/pages/2")
    assert page.status_code == 200 and page.headers["content-type"] == "image/png" and page.content[:4] == b"\x89PNG"
    assert client.get(f"/users/7/sources/{body['original']}/pages/3").status_code == 404
    assert client.get("/users/7/sources/../../secret.pdf/pages/1").status_code == 404


def test_pictures_are_sorted_by_their_layout_without_a_vision_key(client: TestClient, deepseek: FakeDeepSeek, monkeypatch) -> None:
    monkeypatch.delenv("OPENAI_API_KEY", raising=False)
    deepseek.replies = [{"sections": [
        {"id": "s1", "kind": "exercise", "topics": ["Soil testing", "pH"]},
        {"id": "s2", "kind": "text", "topics": ["Watering"]},
    ]}]
    response = client.post("/users/7/sources/classify", json={
        "pictures": [
            {"id": "a", "src": "", "kind": "drawing", "caption": ""},
            {"id": "b", "src": "", "kind": "picture", "caption": "Table 2: Yields"},
            {"id": "c", "src": "", "kind": "picture", "caption": "The score of the opening hymn"},
            {"id": "d", "src": "", "kind": "picture", "caption": "A meadow"},
        ],
        "sections": [{"id": "s1", "title": "Try it yourself", "text": "1. Measure..."}, {"id": "s2", "title": "Rain", "text": "Rain falls"}],
    })
    assert response.status_code == 200, response.text
    body = response.json()
    assert [(p["id"], p["kind"]) for p in body["pictures"]] == [("a", "diagram"), ("b", "table"), ("c", "music"), ("d", "image")]
    assert body["sections"] == [
        {"id": "s1", "kind": "exercise", "topics": ["Soil testing", "pH"]},
        {"id": "s2", "kind": "text", "topics": ["Watering"]},
    ]


def test_pictures_are_looked_at_by_the_vision_model_when_there_is_a_key(client: TestClient, monkeypatch, tmp_path: Path) -> None:
    from phansora.products.talari import library

    monkeypatch.setenv("OPENAI_API_KEY", "test-key")
    monkeypatch.setenv("OPENAI_MODEL", "gpt-5-nano")
    folder = tmp_path / "talari" / "images" / "7"
    folder.mkdir(parents=True)
    (folder / ("a" * 32 + ".png")).write_bytes(png())
    sent: list[dict] = []

    class Completions:
        async def create(self, **kwargs):
            sent.append(kwargs)
            message = type("M", (), {"content": json.dumps({"kind": "music", "description": "Four bars of a waltz."})})
            return type("R", (), {"choices": [type("C", (), {"message": message})]})

    monkeypatch.setattr(library, "_vision_client", lambda: type("Client", (), {"chat": type("Chat", (), {"completions": Completions()})})())
    response = client.post("/users/7/sources/classify", json={"pictures": [
        {"id": "x", "src": f"/talari/users/7/images/{'a' * 32}.png", "kind": "picture", "caption": ""},
        {"id": "y", "src": "/talari/users/7/images/missing.png", "kind": "drawing", "caption": ""},
    ]})
    body = response.json()
    assert body["pictures"] == [
        {"id": "x", "kind": "music", "description": "Four bars of a waltz."},
        {"id": "y", "kind": "diagram", "description": ""},
    ]
    assert len(sent) == 1 and sent[0]["model"] == "gpt-5-nano"
    assert sent[0]["messages"][0]["content"][1]["image_url"]["url"].startswith("data:image/jpeg;base64,")


# --- a picture made by a conversation ---------------------------------------------------

BRIEF = {"book": {"title": "Gate Keeper", "description": "A girl finds a gate in the fog."}, "purpose": "cover", "mode": "background", "look": "Art Nouveau, watercolor"}


def test_a_picture_is_planned_one_question_at_a_time(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"question": {"id": "Subject", "text": "What should the cover show?", "suggestions": ["A gate in fog", "A gate in fog", "Let Talari decide", "A girl with a key", "x" * 200]}, "summary": "A gate in the fog."}]
    response = client.post("/users/7/images/guide", json=BRIEF)
    assert response.status_code == 200, response.text
    body = response.json()
    assert body["question"]["id"] == "subject"
    # No repeats, no "decide" (the app offers that itself), and each answer kept short.
    assert body["question"]["suggestions"] == ["A gate in fog", "A girl with a key", "x" * 80]
    assert body["summary"] == "A gate in the fog."
    user = deepseek.calls[0]["user"]
    assert "Art Nouveau, watercolor" in user and "never ask about art style" in user
    assert "where the title and words will sit" in user


def test_a_question_already_asked_or_full_of_jargon_is_not_put_to_the_author(client: TestClient, deepseek: FakeDeepSeek) -> None:
    answered = {**BRIEF, "answers": [{"id": "subject", "question": "What should the cover show?", "answer": "", "decide": True}]}
    deepseek.replies = [{"question": {"id": "subject", "text": "What else should it show?", "suggestions": []}, "summary": "A gate."}]
    assert client.post("/users/7/images/guide", json=answered).json()["question"] is None
    assert "Let Talari decide" in deepseek.calls[0]["user"]
    deepseek.replies = [{"question": {"id": "ratio", "text": "Which aspect ratio for the prompt?", "suggestions": ["2:3"]}, "summary": "A gate."}]
    assert client.post("/users/7/images/guide", json=answered).json()["question"] is None


def test_after_six_answers_only_the_summary_comes_back(client: TestClient, deepseek: FakeDeepSeek) -> None:
    answers = [{"id": f"q{i}", "question": f"Question {i}?", "answer": "Yes"} for i in range(6)]
    deepseek.replies = [{"question": {"id": "more", "text": "And the colors?", "suggestions": ["Warm"]}, "summary": "A full picture."}]
    body = client.post("/users/7/images/guide", json={**BRIEF, "answers": answers}).json()
    assert body == {"question": None, "summary": "A full picture."}


def test_the_answers_become_a_description_the_author_never_sees(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"prompt": "An iron gate in morning fog, a girl in green before it.", "caption": "The gate", "summary": "A girl at a gate in fog.", "shape": "tall"}]
    response = client.post("/users/7/images/compose", json={**BRIEF, "mode": "media", "purpose": "interior", "pageText": "She stood at the gate."})
    assert response.status_code == 200, response.text
    assert response.json() == {"prompt": "An iron gate in morning fog, a girl in green before it.", "caption": "The gate", "summary": "A girl at a gate in fog.", "shape": "square"}
    user = deepseek.calls[0]["user"]
    assert "She stood at the gate." in user and "Never include words" in user


def test_a_change_rewrites_the_description_and_drawing_again_needs_no_words(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"prompt": "The gate at dusk, darker.", "caption": "Gate at dusk", "summary": "The gate, darker.", "shape": "portrait"}]
    body = client.post("/users/7/images/compose", json={**BRIEF, "previous": "The gate at noon.", "instruction": "Make the lighting darker"}).json()
    assert body["prompt"] == "The gate at dusk, darker."
    assert "The gate at noon." in deepseek.calls[0]["user"] and "Make the lighting darker" in deepseek.calls[0]["user"]
    again = client.post("/users/7/images/compose", json={**BRIEF, "previous": "The gate at noon."}).json()
    assert again["prompt"] == "The gate at noon." and len(deepseek.calls) == 1


def test_background_designs_are_made_for_the_book_and_its_look(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"designs": [
        {"title": "Flowing dragon corner ornaments", "description": "Coiling dragons in each corner.", "kinds": ["corners", "nonsense"], "palette": ["#1f3b2d", "#c9a227", "red"]},
        {"title": "Flowing dragon corner ornaments", "description": "The same again.", "kinds": ["corners"]},
        {"title": "Patterned parchment", "description": "Soft parchment with a faint dragon-scale pattern.", "kinds": ["texture", "pattern"], "palette": []},
        {"title": "Old map", "kinds": []},
        {"title": "The prompt seed", "kinds": ["scene"]},
        {"title": "Ornamental dragon emblem", "description": "A round medallion.", "kinds": ["emblem"]},
    ]}]
    body = {**BRIEF, "book": {"title": "Dragons of the East"}, "size": "6 × 9 in, portrait", "exclude": ["Old map"], "count": 4}
    response = client.post("/users/7/images/designs", json=body)
    assert response.status_code == 200, response.text
    designs = response.json()["designs"]
    assert [d["title"] for d in designs] == ["Flowing dragon corner ornaments", "Patterned parchment", "Ornamental dragon emblem"]
    assert designs[0]["kinds"] == ["corners"] and designs[0]["palette"] == ["#1f3b2d", "#c9a227"]
    user = deepseek.calls[0]["user"]
    assert "Dragons of the East" in user and "Art Nouveau, watercolor" in user and "6 × 9 in, portrait" in user
    assert "corner ornaments" in user and "Old map" in user and "calm area for the title" in user


def test_after_a_design_the_questions_are_about_what_it_needs(client: TestClient, deepseek: FakeDeepSeek) -> None:
    design = {"title": "Flowing dragon corner ornaments", "description": "Coiling dragons in each corner.", "kinds": ["corners"]}
    deepseek.replies = [{"question": {"id": "reach", "text": "How far should the dragons reach in?", "suggestions": ["Just the corners", "A third of the way"]}, "summary": "Dragons in the corners."}]
    client.post("/users/7/images/guide", json={**BRIEF, "designs": [design]})
    user = deepseek.calls[0]["user"]
    assert "THE BACKGROUND DESIGN THE AUTHOR CHOSE" in user and "how far in they reach" in user
    assert "never about the look, the size or shape of the page" in user
    deepseek.replies = [{"prompt": "Four coiling dragons in the corners of a pale ground.", "caption": "Dragon corners", "summary": "Dragons in the corners.", "shape": "portrait"}]
    client.post("/users/7/images/compose", json={**BRIEF, "designs": [design]})
    rules = deepseek.calls[1]["user"]
    assert "ornaments in the corners only" in rules and "Keep the open part of the page plain and calm" in rules
    assert "No figures, creatures or scenery beyond the decoration itself." in rules


def test_a_background_of_a_chosen_type_and_texture_is_planned_without_the_book(client: TestClient, deepseek: FakeDeepSeek) -> None:
    # The app sends no book for a background: it is made from what the author says of it.
    brief = {"purpose": "interior", "mode": "background", "look": "Realism", "backgroundType": "atmospheric", "texture": "stone",
             "answers": [{"id": "background-type", "question": "What type of background?", "answer": "Atmospheric"}]}
    deepseek.replies = [{"question": {"id": "mood", "text": "What mood should it have?", "suggestions": ["Dark and quiet"]}, "summary": "A misty stone ground."}]
    response = client.post("/users/7/images/guide", json=brief)
    assert response.status_code == 200, response.text
    user = deepseek.calls[0]["user"]
    assert "THE BOOK" not in user and "Untitled" not in user
    assert "THE BACKGROUND THE AUTHOR CHOSE" in user and "mist, haze, glow" in user and "stone and marble" in user
    assert "Never ask about the type of background or its texture" in user and "fit this background" in user
    deepseek.replies = [{"prompt": "Mist over a dark stone ground.", "caption": "Misty stone", "summary": "Mist over stone.", "shape": "portrait"}]
    client.post("/users/7/images/compose", json=brief)
    rules = deepseek.calls[1]["user"]
    assert "It is light and mood more than things" in rules and "Its surface is stone and marble." in rules
    assert "No figures, creatures or story scenes in it." in rules and "suits this background best" in rules


def test_a_background_in_the_authors_own_words_is_matched_to_the_set_choices(client: TestClient, deepseek: FakeDeepSeek) -> None:
    types = [{"id": "atmospheric", "label": "Atmospheric"}, {"id": "decorative", "label": "Decorative"}, {"id": "scenic", "label": "Scenic"}]
    textures = [{"id": "stone", "label": "Stone & Marble"}, {"id": "distressed", "label": "Distressed"}, {"id": "wood", "label": "Wood"}]
    deepseek.replies = [{"types": ["atmospheric", "nonsense", "decorative", "scenic"], "textures": ["stone", "stone", "distressed", "glass"]}]
    response = client.post("/users/7/images/background-choices", json={"words": "a stone occult background", "types": types, "textures": textures})
    assert response.status_code == 200, response.text
    # Only choices offered, each once, at most two types and three textures.
    assert response.json() == {"types": ["atmospheric", "decorative"], "textures": ["stone", "distressed"]}
    user = deepseek.calls[0]["user"]
    assert "a stone occult background" in user and "- stone: Stone & Marble" in user and "- scenic: Scenic" in user
    assert client.post("/users/7/images/background-choices", json={"words": "", "types": types}).status_code == 422


REMIX = {**BRIEF, "mode": "media", "purpose": "interior", "look": "Fantasy Art, watercolor", "previous": "A green dragon curled on a rock.", "instruction": "Make this dragon look more real"}


def test_a_remix_asks_a_couple_of_questions_within_the_pictures_look(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"questions": [
        {"id": "Realism", "text": "How realistic should the dragon be?", "suggestions": ["A touch more lifelike", "Detailed and lifelike", "As real as watercolor allows"]},
        {"id": "render", "text": "Which render engine should it use?", "suggestions": ["Octane"]},
        {"id": "scales", "text": "Which parts should look most real?", "suggestions": []},
        {"id": "light", "text": "Should the light change too?", "suggestions": ["Keep it", "Warmer", "Dramatic"]},
        {"id": "more", "text": "Anything else?", "suggestions": ["No"]},
    ]}]
    response = client.post("/users/7/images/remix-questions", json=REMIX)
    assert response.status_code == 200, response.text
    # Two at most; none with machinery words or without answers to pick.
    assert [q["id"] for q in response.json()["questions"]] == ["realism", "light"]
    user = deepseek.calls[0]["user"]
    assert "Make this dragon look more real" in user and "A green dragon curled on a rock." in user
    assert "stays drawn as Fantasy Art, watercolor" in user and "never a photograph" in user
    assert client.post("/users/7/images/remix-questions", json={**REMIX, "instruction": ""}).status_code == 422


def test_a_remix_is_described_with_what_the_author_said_and_in_its_look(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"prompt": "A green dragon with lifelike scales.", "caption": "Dragon", "summary": "A dragon.", "shape": "square"}]
    clarifications = [
        {"question": "How realistic should the dragon be?", "answer": "Detailed and lifelike"},
        {"question": "Should the light change too?", "decide": True},
    ]
    response = client.post("/users/7/images/compose", json={**REMIX, "clarifications": clarifications})
    assert response.status_code == 200, response.text
    user = deepseek.calls[0]["user"]
    assert "WHAT THEY SAID WHEN ASKED ABOUT IT" in user
    assert "- How realistic should the dragon be?\n  Detailed and lifelike" in user
    assert "- Should the light change too?\n  (Let Talari decide.)" in user
    assert "The picture stays drawn as Fantasy Art, watercolor: make the change within that look." in user


def test_with_no_design_yet_a_background_asks_what_kind_it_is(client: TestClient, deepseek: FakeDeepSeek) -> None:
    deepseek.replies = [{"question": {"id": "kind", "text": "What kind of background?", "suggestions": ["A border"]}, "summary": ""}]
    client.post("/users/7/images/guide", json=BRIEF)
    assert "a decorative border, corner ornaments, a texture" in deepseek.calls[0]["user"]
