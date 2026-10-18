"""What an upload holds beyond its running text (POST /sources/extract).

sources.py reads a document into chapters. This module reads the same file again for what
that leaves out, keeping it as it is rather than as a summary of it:

  * its structure — the sections under each chapter, by the document's own headings: a
    PDF's deeper outline entries or, without them, its larger and bolder type; a DOCX's
    Heading 2 and 3; an EPUB's h2–h4. For a PDF with no outline and no "Chapter 3" lines,
    the largest headings that recur become the chapters, so "Part 1, Part 2" is the last
    resort rather than the usual case.
  * its pictures (raster images, kept at their own resolution), its drawings (vector
    diagrams, rendered from the page) and its tables (their cells as text, plus a picture
    of the table as it was printed), each with its page and a caption found beside it.

Pictures are saved where Talari keeps every picture of the user's (``images/<user_id>/``),
named by their content, so the same picture imported twice is one file, and the app's
picture proxy and its PDF export already know how to fetch them.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
import time
from dataclasses import dataclass, field
from pathlib import Path

from phansora.products.book_alchemy.parsers import classify_block

logger = logging.getLogger("talari")

MAX_ASSETS = 400
# Finding a PDF's pictures and tables stops after this long, keeping what was found. The
# upload does not wait for it: the app asks for them once the import is answered
# (POST …/sources/{name}/assets), since a long illustrated book takes minutes.
ASSET_SECONDS = 600
# Past this many vector paths a page is a chart or a pattern, measured as one region.
MAX_PATHS = 2000
# Past this many, a page is music notation, a chart or a pattern, not a ruled table: finding
# tables by their ruling would take seconds a page there, for nothing.
TABLE_MAX_PATHS = 400
# A picture covering this much of its page, on this share of the pages, makes the PDF a
# scan: its text layer was read by OCR, so the size of its type says nothing about headings,
# and each page's picture is the page itself, not a picture in it.
SCAN_COVER = 0.85
SCAN_PAGES = 0.6
# How many pages, spread through the book, are looked at to tell a scan.
SCAN_SAMPLE = 16
# Chapters found by the size of the type are believed only when they look like chapters:
# no more than one every few pages, and long enough to read. Otherwise the parts are the
# better guess. Sections the same, more loosely.
PAGES_PER_CHAPTER = 3
CHAPTER_MEDIAN_WORDS = 150
SECTION_MEDIAN_WORDS = 25
# Smaller than this (either side, in pixels) is an icon, a bullet or a rule, not a picture.
MIN_PICTURE_PX = 80
# Larger than this is scaled down: past it a picture only adds bytes, never sharpness on a page.
MAX_PICTURE_PX = 2400
# A picture on this many pages is a logo, a page background or an ornament.
REPEATED_ON_PAGES = 3
# A region of a page (a drawing, a table) is rendered at 3× — 216 dpi.
RENDER_ZOOM = 3
# A drawing is this many vector paths at least, over this share of the page at least.
DRAWING_MIN_PATHS = 8
DRAWING_MIN_AREA = 0.03
MAX_TABLE_ROWS = 300
MAX_TABLE_COLS = 20
MAX_CELL_CHARS = 400
# A heading is short; a line of body text in bold is not one.
HEADING_MAX_WORDS = 14
HEADING_MAX_CHARS = 120
# Type this much larger than the body's is a heading.
HEADING_SCALE = 1.15
CAPTION = re.compile(
    r"^(?:fig(?:ure)?\.?|table|diagram|chart|graph|plate|illustration|map|example|exercise|score|chart)\s*[\dIVX]*\b",
    re.IGNORECASE,
)
CAPTION_MAX_WORDS = 40
# How far from a picture, in points, its caption may be.
CAPTION_REACH = 48


@dataclass
class Para:
    text: str
    page: int | None = None
    y: float = 0.0


@dataclass
class Heading:
    level: int
    title: str
    page: int | None = None
    y: float = 0.0


@dataclass
class Section:
    chapter_index: int
    level: int
    title: str
    page_start: int | None
    page_end: int | None
    text: str


@dataclass
class Asset:
    kind: str                      # picture | drawing | table
    src: str | None
    width: int | None
    height: int | None
    page: int | None
    caption: str = ""
    rows: list[list[str]] | None = None
    header: bool = False
    chapter_index: int = 0
    section_index: int | None = None
    # Where it sits, to place it under a chapter and a section: (page, y) for a PDF, the
    # running order of the document for the rest.
    at: tuple[float, float] = (0.0, 0.0)


@dataclass
class StructuredChapter:
    """A chapter read from the document's own headings, for when sources.py found none."""

    title: str
    paragraphs: list[str] = field(default_factory=list)
    pages: list[int] = field(default_factory=list)


@dataclass
class Material:
    sections: list[Section] = field(default_factory=list)
    assets: list[Asset] = field(default_factory=list)


def _clean(text: str) -> str:
    return " ".join((text or "").split())


# --- pictures -------------------------------------------------------------------------


class Pictures:
    """Where a document's pictures are saved, and how: normalized, scaled, named by content."""

    def __init__(self, folder: Path, user_id: str) -> None:
        self.folder = folder
        self.user_id = user_id

    def save(self, data: bytes) -> tuple[str, int, int] | None:
        from PIL import Image

        try:
            image = Image.open(io.BytesIO(data))
            image.load()
        except Exception:  # noqa: BLE001 — a picture Pillow cannot read is left out, not an error
            return None
        if min(image.size) < MIN_PICTURE_PX:
            return None
        photo = image.format == "JPEG" and image.mode in ("RGB", "L")
        if image.mode not in ("RGB", "RGBA", "L", "LA"):
            image = image.convert("RGBA" if "A" in image.getbands() else "RGB")
        if max(image.size) > MAX_PICTURE_PX:
            image.thumbnail((MAX_PICTURE_PX, MAX_PICTURE_PX))
        out = io.BytesIO()
        if photo:
            image.convert("RGB").save(out, "JPEG", quality=90)
            suffix = "jpg"
        else:
            image.save(out, "PNG", optimize=True)
            suffix = "png"
        payload = out.getvalue()
        name = f"{hashlib.sha256(payload).hexdigest()[:32]}.{suffix}"
        self.folder.mkdir(parents=True, exist_ok=True)
        path = self.folder / name
        if not path.exists():
            path.write_bytes(payload)
        return f"/talari/users/{self.user_id}/images/{name}", image.size[0], image.size[1]


def _rows(rows: list[list[object]]) -> list[list[str]]:
    """A table's cells as plain text, trimmed to a size a page can hold."""
    out = [[_clean(str(cell or ""))[:MAX_CELL_CHARS] for cell in row[:MAX_TABLE_COLS]] for row in rows[:MAX_TABLE_ROWS]]
    width = max((len(row) for row in out), default=0)
    return [row + [""] * (width - len(row)) for row in out if any(row)]


# --- PDF ------------------------------------------------------------------------------


def _is_bold(span: dict) -> bool:
    return bool(span.get("flags", 0) & 16) or "bold" in str(span.get("font", "")).lower()


def _pdf_blocks(pdf) -> list[tuple[int, dict]]:
    """Every text block of every page, as (page number from 1, block), in reading order."""
    out = []
    for pno in range(pdf.page_count):
        page = pdf.load_page(pno)
        for block in page.get_text("dict", sort=True).get("blocks", []):
            if block.get("type") == 0 and block.get("lines"):
                out.append((pno + 1, block))
    return out


def _block_text(block: dict) -> str:
    return "\n".join(" ".join(span.get("text", "") for span in line.get("spans", [])).strip() for line in block["lines"]).strip()


def _body_size(blocks: list[tuple[int, dict]]) -> float:
    """The size most of the text is set in: the body, against which headings stand out."""
    counts: dict[float, int] = {}
    for _, block in blocks:
        for line in block["lines"]:
            for span in line.get("spans", []):
                size = round(float(span.get("size", 0)) * 2) / 2
                counts[size] = counts.get(size, 0) + len(span.get("text", "").strip())
    return max(counts, key=counts.get) if counts else 10.0


def _running_keys(blocks: list[tuple[int, dict]], pages: int) -> set[str]:
    """Lines repeated at the top or the foot of many pages: running heads and folios."""
    if pages < 4:
        return set()
    by_page: dict[int, list[str]] = {}
    for pno, block in blocks:
        by_page.setdefault(pno, []).append(_block_text(block))
    counts: dict[str, int] = {}
    for texts in by_page.values():
        for text in {texts[0], texts[-1]}:
            if 0 < len(text) <= 80:
                key = re.sub(r"\d+", "#", _clean(text))
                counts[key] = counts.get(key, 0) + 1
    return {key for key, n in counts.items() if n >= max(3, pages // 3)}


def _reads_as_heading(text: str) -> bool:
    """Words, not the noise a text layer makes of music, ornaments or a page number."""
    chars = [c for c in text if not c.isspace()]
    letters = sum(c.isalpha() for c in chars)
    return letters >= 3 and letters >= 0.7 * len(chars) and bool(re.search(r"[^\W\d_]{3,}", text))


def _scanned(pdf) -> bool:
    """Most of a sample of its pages a single picture of the page (see SCAN_COVER). By the
    pictures' boxes alone: naming each picture's xref (get_image_rects) hashes the picture,
    seconds a page for a scan."""
    import fitz

    count = pdf.page_count
    if not count:
        return False
    sample = sorted({round(i * (count - 1) / max(1, SCAN_SAMPLE - 1)) for i in range(min(SCAN_SAMPLE, count))})
    covered = 0
    for pno in sample:
        page = pdf.load_page(pno)
        area = page.rect.width * page.rect.height
        if any(fitz.Rect(info["bbox"]).get_area() >= SCAN_COVER * area for info in page.get_image_info()):
            covered += 1
    return covered >= SCAN_PAGES * len(sample)


class PdfReading:
    """What the chapters (pdf_chapters) and the library (read) both need from one PDF, read
    once: whether it is a scan, its outline, and, only when asked for, its headings and
    paragraphs by the size of their type (a fifth of a second a page on a scan's text layer)."""

    def __init__(self, path: Path) -> None:
        import fitz

        self.path = path
        with fitz.open(path) as pdf:
            self.pages = pdf.page_count
            self.scanned = _scanned(pdf)
            self.toc = pdf.get_toc(simple=True)
        self._nodes: tuple[list[Heading], list[Para]] | None = None

    def nodes(self) -> tuple[list[Heading], list[Para]]:
        if self._nodes is None:
            import fitz

            with fitz.open(self.path) as pdf:
                self._nodes = _pdf_nodes(pdf)
        return self._nodes


def _median(values: list[int]) -> int:
    ordered = sorted(values)
    return ordered[len(ordered) // 2] if ordered else 0


def _pdf_nodes(pdf) -> tuple[list[Heading], list[Para]]:
    """A PDF's headings, by the size and weight of their type, and its paragraphs."""
    blocks = _pdf_blocks(pdf)
    body = _body_size(blocks)
    running = _running_keys(blocks, pdf.page_count)
    headings: list[tuple[float, Heading]] = []
    paras: list[Para] = []
    for pno, block in blocks:
        text = _block_text(block)
        if not text or re.sub(r"\d+", "#", _clean(text)) in running or re.fullmatch(r"[\divxlcdm\s.-]+", text, re.I):
            continue
        # A block can hold a heading and a smaller line beside it (a date, a byline): the
        # heading is only the lines in its own size, and the rest is text.
        lines = [line for line in block["lines"] if any(s.get("text", "").strip() for s in line.get("spans", []))]
        sizes = [max(float(s.get("size", 0)) for s in line.get("spans", []) if s.get("text", "").strip()) for line in lines]
        size = max(sizes, default=body)
        top = [line for line, s in zip(lines, sizes) if s >= size - 0.5]
        rest = [line for line, s in zip(lines, sizes) if s < size - 0.5]
        spans = [span for line in top for span in line.get("spans", []) if span.get("text", "").strip()]
        bold = bool(spans) and all(_is_bold(span) for span in spans)
        flat = _clean(" ".join(" ".join(s.get("text", "") for s in line.get("spans", [])) for line in top))
        short = len(top) <= 2 and len(flat.split()) <= HEADING_MAX_WORDS and len(flat) <= HEADING_MAX_CHARS
        if flat and short and not re.search(r"[.,;:]$", flat) and _reads_as_heading(flat) and (size >= body * HEADING_SCALE or (bold and size >= body * 0.98)):
            headings.append((round(size * 2) / 2 + (0.25 if bold else 0), Heading(level=0, title=flat, page=pno, y=block["bbox"][1])))
            text = "\n".join(" ".join(s.get("text", "") for s in line.get("spans", [])).strip() for line in rest)
            if text and classify_block(text) != "apparatus":
                paras.append(Para(text=_clean(text), page=pno, y=block["bbox"][1] + 0.1))
        elif classify_block(text) != "apparatus":
            paras.append(Para(text=_clean(text), page=pno, y=block["bbox"][1]))
    # Levels by size: the largest heading type is level 1, the next 2, the rest 3.
    sizes = sorted({size for size, _ in headings}, reverse=True)
    for size, heading in headings:
        heading.level = min(3, sizes.index(size) + 1)
    return [heading for _, heading in headings], paras


def _pos(page: int | None, y: float) -> tuple[float, float]:
    return (float(page or 0), y)


def _structured(headings: list[Heading], paras: list[Para], chapter_level: int) -> list[tuple[Heading, list[Heading], list[Para]]]:
    """Chapters at ``chapter_level``, each with the headings and paragraphs that follow it."""
    out: list[tuple[Heading, list[Heading], list[Para]]] = []
    marks = sorted(headings, key=lambda h: _pos(h.page, h.y))
    starts = [h for h in marks if h.level == chapter_level]
    for i, start in enumerate(starts):
        end = _pos(starts[i + 1].page, starts[i + 1].y) if i + 1 < len(starts) else (float("inf"), 0.0)
        begin = _pos(start.page, start.y)
        inside = [h for h in marks if h.level > chapter_level and begin < _pos(h.page, h.y) < end]
        text = [p for p in paras if begin <= _pos(p.page, p.y) < end]
        out.append((start, inside, text))
    return out


def pdf_chapters(path: Path, reading: PdfReading | None = None) -> list[StructuredChapter]:
    """For a PDF with no outline and no "Chapter 3" lines: chapters by its largest recurring
    headings. Empty when there are fewer than two, or for a scan."""
    try:
        reading = reading or PdfReading(path)
        if reading.scanned:
            return []
        pages = reading.pages
        headings, paras = reading.nodes()
    except Exception:  # noqa: BLE001 — no structure is a fallback, never a failure
        logger.exception("Talari: could not read a PDF's headings")
        return []
    for level in (1, 2, 3):
        if sum(1 for h in headings if h.level == level) >= 2:
            chapters = []
            for start, _, text in _structured(headings, paras, level):
                chapter = StructuredChapter(title=start.title)
                for para in text:
                    chapter.paragraphs.append(para.text)
                    if para.page:
                        chapter.pages.append(para.page)
                if chapter.paragraphs:
                    chapters.append(chapter)
            plausible = (
                2 <= len(chapters) <= max(12, pages // PAGES_PER_CHAPTER)
                and _median([sum(len(p.split()) for p in c.paragraphs) for c in chapters]) >= CHAPTER_MEDIAN_WORDS
            )
            if not plausible:
                logger.info("Talari: %d chapters by type size in %d pages do not look like chapters", len(chapters), pages)
            return chapters if plausible else []
    return []


def _caption_near(page, rect) -> str:
    """The caption of a picture or table: a line beginning "Figure", "Table"… just below or
    above it, else a short line right below it."""
    below, above, near = None, None, None
    for block in page.get_text("blocks", sort=True):
        x0, y0, x1, y1, text = block[0], block[1], block[2], block[3], _clean(block[4])
        if not text or len(text.split()) > CAPTION_MAX_WORDS or x1 < rect.x0 or x0 > rect.x1:
            continue
        if 0 <= y0 - rect.y1 <= CAPTION_REACH:
            if CAPTION.match(text) and below is None:
                below = text
            elif near is None:
                near = text
        elif 0 <= rect.y0 - y1 <= CAPTION_REACH and CAPTION.match(text):
            above = text
    return below or above or near or ""


def _drawing_regions(page, drawings: list) -> list:
    """Clusters of vector paths big enough to be a diagram, not a rule or a box."""
    import fitz

    rects = [fitz.Rect(path["rect"]) for path in drawings if path.get("rect")]
    rects = [r for r in rects if r.width > 1 or r.height > 1]
    if len(rects) > MAX_PATHS:
        whole = fitz.Rect(rects[0])
        for rect in rects[1:]:
            whole |= rect
        rects = [whole] * DRAWING_MIN_PATHS
    clusters: list[list] = []
    for rect in rects:
        grown = rect + (-6, -6, 6, 6)
        hit = [c for c in clusters if c[0].intersects(grown)]
        merged = [fitz.Rect(rect), 1]
        for c in hit:
            merged[0] |= c[0]
            merged[1] += c[1]
            clusters.remove(c)
        clusters.append(merged)
    area = page.rect.width * page.rect.height
    out = []
    for rect, count in clusters:
        share = (rect.width * rect.height) / area
        # A frame round the whole page is a border, however many strokes it has.
        if count >= DRAWING_MIN_PATHS and share >= DRAWING_MIN_AREA and share < 0.9:
            out.append(rect)
    return out


def _pdf_assets(pdf, pictures: Pictures, *, budget: float = ASSET_SECONDS, scanned: bool = False) -> list[Asset]:
    import fitz

    uses: dict[int, int] = {}
    for pno in range(pdf.page_count):
        for xref in {info[0] for info in pdf.load_page(pno).get_images(full=True)}:
            uses[xref] = uses.get(xref, 0) + 1
    assets: list[Asset] = []
    deadline = time.monotonic() + budget
    for pno in range(pdf.page_count):
        if len(assets) >= MAX_ASSETS:
            break
        if time.monotonic() > deadline:
            logger.warning("Talari: stopped looking for pictures and tables at page %d of %d", pno + 1, pdf.page_count)
            break
        page = pdf.load_page(pno)
        zoom = fitz.Matrix(RENDER_ZOOM, RENDER_ZOOM)
        drawings = page.get_drawings()
        tables = []
        # A table is found by its ruling lines, so a page with no vector strokes has none to
        # find, and one with thousands is notation or a pattern (TABLE_MAX_PATHS).
        if 0 < len(drawings) <= TABLE_MAX_PATHS:
            try:
                tables = page.find_tables().tables
            except Exception:  # noqa: BLE001 — an older PyMuPDF, or a page it cannot read for tables
                tables = []
        for table in tables:
            rows = _rows(table.extract() or [])
            if len(rows) < 2 or len(rows[0]) < 2:
                continue
            rect = fitz.Rect(table.bbox)
            shot = pictures.save(page.get_pixmap(matrix=zoom, clip=rect + (-2, -2, 2, 2), alpha=False).tobytes("png"))
            assets.append(Asset(
                kind="table", src=shot[0] if shot else None, width=shot[1] if shot else None, height=shot[2] if shot else None,
                page=pno + 1, caption=_caption_near(page, rect), rows=rows, header=True, at=(pno + 1, rect.y0),
            ))
        table_rects = [fitz.Rect(t.bbox) for t in tables]
        # A scan's pictures are its pages. Elsewhere, where each picture sits, asked once a page:
        # naming the xrefs hashes every picture on it.
        boxes: dict[int, object] = {}
        if not scanned:
            for placed in page.get_image_info(xrefs=True):
                boxes.setdefault(placed.get("xref"), fitz.Rect(placed["bbox"]))
        for info in [] if scanned else page.get_images(full=True):
            xref = info[0]
            if uses.get(xref, 0) >= REPEATED_ON_PAGES:
                continue
            rect = boxes.get(xref)
            if rect is None or rect.width < 36 or rect.height < 36:
                continue
            try:
                raw = pdf.extract_image(xref)
            except Exception:  # noqa: BLE001
                continue
            # A picture with a transparency mask, or in a form Pillow cannot read, is taken
            # as it appears on the page instead.
            data = raw.get("image") if raw and not raw.get("smask") and raw.get("ext") in ("jpeg", "jpg", "png") else None
            if data is None:
                data = page.get_pixmap(matrix=zoom, clip=rect, alpha=False).tobytes("png")
            saved = pictures.save(data)
            if saved:
                assets.append(Asset(kind="picture", src=saved[0], width=saved[1], height=saved[2], page=pno + 1,
                                    caption=_caption_near(page, rect), at=(pno + 1, rect.y0)))
        for rect in _drawing_regions(page, drawings):
            if any(rect.intersects(t) and (rect & t).get_area() > 0.5 * rect.get_area() for t in table_rects):
                continue
            saved = pictures.save(page.get_pixmap(matrix=zoom, clip=rect + (-4, -4, 4, 4), alpha=False).tobytes("png"))
            if saved:
                assets.append(Asset(kind="drawing", src=saved[0], width=saved[1], height=saved[2], page=pno + 1,
                                    caption=_caption_near(page, rect), at=(pno + 1, rect.y0)))
    return assets[:MAX_ASSETS]


def _pdf_material(
    path: Path, chapters: list, outline: list[tuple[int, str]], pictures: Pictures, *, assets: bool, budget: float,
    reading: PdfReading | None = None,
) -> Material:
    import fitz

    reading = reading or PdfReading(path)
    pages, scanned = reading.pages, reading.scanned
    # Sections: the outline's entries below its chapters when it has them, else the type's
    # headings — those smaller than the chapters' own. A scan's type says nothing, so it has
    # sections only from an outline, and its text is not read for headings at all without one.
    chapter_titles = {_clean(c.title).lower() for c in chapters}
    deeper = [
        Heading(level=min(3, int(level)), title=_clean(str(title)), page=int(page), y=0.0)
        for level, title, page in reading.toc
        if page and _clean(str(title)).lower() not in chapter_titles
    ] if outline else []
    headings, paras = ([], []) if scanned and not deeper else reading.nodes()
    if scanned:
        headings = []
    found = []
    if assets:
        with fitz.open(path) as pdf:
            found = _pdf_assets(pdf, pictures, budget=budget, scanned=scanned)
    marks = [h for h in (deeper or headings) if _clean(h.title).lower() not in chapter_titles]
    material = Material(assets=found)
    # Where each chapter begins: at its own heading when the type gave it one, else at the
    # top of its first page. Chapters found by type can share a page, so a heading belongs to
    # the last chapter begun before it, and a section ends where the next chapter begins.
    starts: list[tuple[float, float] | None] = []
    for chapter in chapters:
        if chapter.page_start is None:
            starts.append(None)
            continue
        title = _clean(chapter.title).lower()
        own = next((h for h in headings if h.page == chapter.page_start and _clean(h.title).lower() == title), None)
        starts.append(_pos(chapter.page_start, own.y if own else 0.0))
    owner: dict[int, int] = {}
    for n, heading in enumerate(marks):
        at = _pos(heading.page, heading.y)
        for index, chapter in enumerate(chapters):
            if heading.page and starts[index] is not None and starts[index] <= at and heading.page <= (chapter.page_end or heading.page):
                owner[n] = index
    for index, chapter in enumerate(chapters):
        start, end = chapter.page_start, chapter.page_end
        if start is None:
            continue
        following = [st for st in starts[index + 1:] if st is not None]
        limit = min([(float(end) + 1, 0.0), *following])
        inside = sorted((h for n, h in enumerate(marks) if owner.get(n) == index), key=lambda h: _pos(h.page, h.y))
        for i, heading in enumerate(inside):
            stop = min(_pos(inside[i + 1].page, inside[i + 1].y), limit) if i + 1 < len(inside) else limit
            text = [p for p in paras if _pos(heading.page, heading.y) <= _pos(p.page, p.y) < stop]
            if not text:
                continue
            material.sections.append(Section(
                chapter_index=index, level=heading.level, title=heading.title,
                page_start=text[0].page, page_end=text[-1].page, text="\n\n".join(p.text for p in text),
            ))
    # Headings by type size that cut the text into scraps, or into more pieces than the book
    # has pages twice over, are not its sections.
    if not deeper and material.sections and (
        len(material.sections) > max(20, 2 * pages) or _median([len(s.text.split()) for s in material.sections]) < SECTION_MEDIAN_WORDS
    ):
        logger.info("Talari: %d sections by type size in %d pages do not look like sections", len(material.sections), pages)
        material.sections = []
    _place_by_page(material, chapters)
    return material


def _place_by_page(material: Material, chapters: list) -> None:
    """Each picture and table under the chapter, and the section, its page falls in."""
    for asset in material.assets:
        page = asset.page or 0
        for index, chapter in enumerate(chapters):
            if chapter.page_start is not None and chapter.page_start <= page <= (chapter.page_end or page):
                asset.chapter_index = index
        best = None
        for index, section in enumerate(material.sections):
            if section.chapter_index == asset.chapter_index and section.page_start is not None and section.page_start <= page:
                best = index
        asset.section_index = best


# --- DOCX ------------------------------------------------------------------------------


def _docx_material(path: Path, pictures: Pictures) -> tuple[list[StructuredChapter], list[tuple[int, str, int, list[str]]], list[tuple[int, Asset]]]:
    """A DOCX in its own order: headings by level, paragraphs, tables and pictures. Returns
    the chapters by its top heading level, the sections under them as (title, level,
    paragraphs) per chapter in order, and the assets with the chapter they fall in."""
    import docx
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    nodes: list[tuple] = []
    body = document.element.body
    for element in body.iterchildren():
        if element.tag == qn("w:p"):
            paragraph = Paragraph(element, document)
            style = (paragraph.style.name if paragraph.style is not None else "") or ""
            text = _clean(paragraph.text)
            level = None
            match = re.match(r"heading\s*(\d)", style, re.IGNORECASE)
            if match:
                level = int(match.group(1))
            elif style.lower() == "title":
                level = 1
            for blip in element.iter(qn("a:blip")):
                rid = blip.get(qn("r:embed"))
                part = document.part.related_parts.get(rid) if rid else None
                saved = pictures.save(part.blob) if part is not None else None
                if saved:
                    nodes.append(("picture", Asset(kind="picture", src=saved[0], width=saved[1], height=saved[2], page=None)))
            if not text:
                continue
            if level:
                nodes.append(("heading", level, text))
            elif style.lower() == "caption" and nodes and nodes[-1][0] in ("picture", "table"):
                nodes[-1][1].caption = text
            else:
                nodes.append(("para", text))
        elif element.tag == qn("w:tbl"):
            table = Table(element, document)
            rows = _rows([[cell.text for cell in row.cells] for row in table.rows])
            if len(rows) >= 2 and len(rows[0]) >= 2:
                nodes.append(("table", Asset(kind="table", src=None, width=None, height=None, page=None, rows=rows, header=True)))

    levels = sorted({node[1] for node in nodes if node[0] == "heading"})
    top = next((lvl for lvl in levels if sum(1 for n in nodes if n[0] == "heading" and n[1] == lvl) >= 2), None)
    chapters: list[StructuredChapter] = []
    sections: list[tuple[int, str, int, list[str]]] = []
    assets: list[tuple[int, Asset]] = []
    current: StructuredChapter | None = None
    for node in nodes:
        if node[0] == "heading" and top is not None and node[1] <= top:
            current = StructuredChapter(title=node[2])
            chapters.append(current)
        elif node[0] == "heading":
            if current is None:
                current = StructuredChapter(title=node[2])
                chapters.append(current)
            else:
                sections.append((len(chapters) - 1, node[2], node[1], []))
        elif node[0] == "para":
            if current is None:
                current = StructuredChapter(title="")
                chapters.append(current)
            current.paragraphs.append(node[1])
            if sections and sections[-1][0] == len(chapters) - 1:
                sections[-1][3].append(node[1])
        else:
            assets.append((max(0, len(chapters) - 1), node[1]))
    return [c for c in chapters if c.paragraphs], sections, assets


# --- EPUB ------------------------------------------------------------------------------


def _epub_material(path: Path, chapters: list, pictures: Pictures) -> Material:
    """An EPUB's sections (h2–h4 after each document's first heading), its pictures (with
    their alt text or figure caption) and its tables, placed under the chapter whose
    document they are in."""
    import posixpath

    import ebooklib
    from bs4 import BeautifulSoup
    from ebooklib import epub

    from phansora.products.book_alchemy.parsers import _is_epub_nav

    book = epub.read_epub(str(path))
    titles = [_clean(c.title).lower() for c in chapters]
    material = Material()
    chapter_index = 0
    for item in book.get_items_of_type(ebooklib.ITEM_DOCUMENT):
        soup = BeautifulSoup(item.get_content(), "html.parser")
        if _is_epub_nav(item, soup):
            continue
        first = soup.find(["h1", "h2", "h3"])
        if first is not None:
            label = _clean(first.get_text(" ", strip=True)).lower()
            if label in titles[chapter_index:]:
                chapter_index = titles.index(label, chapter_index)
        current: tuple[str, int, list[str]] | None = None
        for element in soup.find_all(["h2", "h3", "h4", "p", "li", "blockquote", "img", "table"]):
            if element.name in ("h2", "h3", "h4") and element is not first:
                if current and current[2]:
                    material.sections.append(Section(chapter_index, current[1], current[0], None, None, "\n\n".join(current[2])))
                current = (_clean(element.get_text(" ", strip=True)), int(element.name[1]) - 1, [])
            elif element.name in ("p", "li", "blockquote") and current is not None and not element.find_parent("table"):
                text = _clean(element.get_text(" ", strip=True))
                if text:
                    current[2].append(text)
            elif element.name == "img" and len(material.assets) < MAX_ASSETS:
                src = element.get("src") or ""
                href = posixpath.normpath(posixpath.join(posixpath.dirname(item.get_name()), src.split("#")[0]))
                target = book.get_item_with_href(href)
                saved = pictures.save(target.get_content()) if target is not None else None
                if saved:
                    figure = element.find_parent("figure")
                    caption = figure.find("figcaption") if figure is not None else None
                    text = _clean(caption.get_text(" ", strip=True)) if caption is not None else _clean(element.get("alt") or "")
                    material.assets.append(Asset(kind="picture", src=saved[0], width=saved[1], height=saved[2], page=None,
                                                 caption=text, chapter_index=chapter_index,
                                                 section_index=len(material.sections) if current else None))
            elif element.name == "table" and len(material.assets) < MAX_ASSETS:
                trs = element.find_all("tr")
                rows = _rows([[cell.get_text(" ", strip=True) for cell in tr.find_all(["th", "td"])] for tr in trs])
                if len(rows) >= 2 and len(rows[0]) >= 2:
                    caption = element.find("caption")
                    material.assets.append(Asset(
                        kind="table", src=None, width=None, height=None, page=None, rows=rows,
                        header=bool(trs and trs[0].find("th")), caption=_clean(caption.get_text(" ", strip=True)) if caption else "",
                        chapter_index=chapter_index, section_index=len(material.sections) if current else None,
                    ))
        if current and current[2]:
            material.sections.append(Section(chapter_index, current[1], current[0], None, None, "\n\n".join(current[2])))
    return material


# --- the entry point -------------------------------------------------------------------


def read(
    path: Path, fmt: str, chapters: list, *, outline: list, pictures: Pictures, docx_sections=None, docx_assets=None,
    assets: bool = True, budget: float = ASSET_SECONDS, reading: PdfReading | None = None,
) -> Material:
    """The sections and the assets of a document whose chapters sources.py has read. Without
    ``assets``, a PDF's pictures and tables are left to be looked for later — the slow part;
    a DOCX's and an EPUB's are only read out of the file, and come with it."""
    try:
        if fmt == "pdf":
            return _pdf_material(path, chapters, outline, pictures, assets=assets, budget=budget, reading=reading)
        if fmt == "epub":
            return _epub_material(path, chapters, pictures)
        if fmt == "docx":
            material = Material()
            for chapter_index, title, level, paragraphs in docx_sections or []:
                if paragraphs:
                    material.sections.append(Section(chapter_index, max(1, level - 1), title, None, None, "\n\n".join(paragraphs)))
            for chapter_index, asset in docx_assets or []:
                asset.chapter_index = chapter_index
                material.assets.append(asset)
            return material
    except Exception:  # noqa: BLE001 — the text is the import; what is found beside it is a bonus
        logger.exception("Talari: could not read the structure and pictures of a %s upload", fmt)
    return Material()


def docx_structure(path: Path, pictures: Pictures):
    """A DOCX's chapters by its top heading level, with its sections and assets (see
    _docx_material), or None when it cannot be read that way."""
    try:
        return _docx_material(path, pictures)
    except Exception:  # noqa: BLE001
        logger.exception("Talari: could not read a DOCX's headings")
        return None
